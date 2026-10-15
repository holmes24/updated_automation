"""Create a formatted, date-wise attendance Excel workbook."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MISSING = "Not Assigned"
NOT_PROVIDED = "Not Provided"
NO_RESULT = "—"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
TITLE_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(bold=True, color="FFFFFF")

THIN = Side(style="thin", color="D9E2F3")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MARKER_FILLS = {
    "P": PatternFill("solid", fgColor="C6EFCE"),
    "L": PatternFill("solid", fgColor="D9EAF7"),
    "AL": PatternFill("solid", fgColor="FF9999"),
    "HD": PatternFill("solid", fgColor="FFE4B5"),
    "WO": PatternFill("solid", fgColor="D9D9D9"),
    "H": PatternFill("solid", fgColor="D9EAD3"),
    "A": PatternFill("solid", fgColor="FFC7CE"),
    "F": PatternFill("solid", fgColor="FFEB9C"),
    "M": PatternFill("solid", fgColor="C6EFCE"),
    "N/A": PatternFill("solid", fgColor="D9D9D9"),
    NO_RESULT: PatternFill("solid", fgColor="F2F2F2"),
}


def _value(value, fallback=MISSING):
    if value is None:
        return fallback
    if isinstance(value, str) and not value.strip():
        return fallback
    return value


def _style_sheet(ws, header_row: int = 1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{header_row + 1}"
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.row_dimensions[header_row].height = 28


def _fit_columns(ws, max_width: int = 28):
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=10,
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 11), max_width)


def build_export_workbook(
    coordinators: list[dict],
    employees: list[dict],
    grids: dict[str, dict[str, str] | None],
    days: list[str],
) -> bytes:
    wb = Workbook()
    people = [
        {
            "employee_id": pc["pc_id"],
            "name": pc.get("name"),
        }
        for pc in coordinators
    ] + employees

    # --- Summary Sheet ---
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False

    summary.append(["MONTHLY ATTENDANCE REPORT"])
    summary.merge_cells("A1:B1")
    summary["A1"].fill = TITLE_FILL
    summary["A1"].font = Font(bold=True, size=15, color="FFFFFF")
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 30

    summary.append(["Report Start", days[0] if days else NOT_PROVIDED])
    summary.append(["Report End", days[-1] if days else NOT_PROVIDED])
    summary.append([
        "Total Attendance Participants",
        sum(grids.get(str(person["employee_id"])) is not None for person in people),
    ])
    summary.append([])

    summary.append(["Code", "Meaning"])
    legend_row = summary.max_row
    for cell in summary[legend_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    legend = [
        ("P", "Present — on time"),
        ("L", "Present — late"),
        ("A", "Absent"),
        ("AL", "Absent — after grace period"),
        ("HD", "Half day"),
        ("WO", "Week off"),
        ("H", "Holiday"),
        ("F", "Flagged — awaiting review"),
        ("M", "Manual attendance entry"),
        ("N/A", "Inactive on this date"),
        (NO_RESULT, "No result yet — today or future date"),
    ]
    for marker, meaning in legend:
        summary.append([marker, meaning])
        summary.cell(summary.max_row, 1).fill = MARKER_FILLS.get(marker, PatternFill())

    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 50

    # --- Department attendance worksheets ---
    roles = [
        ("FS", "Field Supervisors"),
        ("LS", "Local supervisors"),
        ("LSS", "Local supervisor scanning"),
    ]
    person_by_id = {str(person["employee_id"]): person for person in people}
    for role, sheet_title in roles:
        ws_pc = wb.create_sheet(sheet_title)
        # Editable daily codes; totals recalculate automatically in Excel.
        headers = ["ID", "Location", "Name", "Reports To", "Position"] + days + ["Present", "Absent", "HD", "WO", "Holidays", "Late", "AL"]
        ws_pc.append(headers)
        from openpyxl.worksheet.datavalidation import DataValidation
        validation = DataValidation(type="list", formula1='"P,L,A,AL,HD,WO,H,M,F,N/A"', allow_blank=True)
        validation.error = "Select an attendance code from the list"
        validation.showErrorMessage = True
        ws_pc.add_data_validation(validation)
        role_people = [
            person for person in people
            if person.get("role", "FS") == role
            and grids.get(str(person.get("employee_id", ""))) is not None
        ]
        for emp in sorted(role_people, key=lambda x: str(x.get("name") or "").casefold()):
            emp_id = str(emp.get("employee_id", ""))
            grid = grids.get(emp_id) or {}
            markers = [_value(grid.get(day), NO_RESULT) for day in days]
            site = emp.get("assigned_site") or {}
            location = site.get("name") if isinstance(site, dict) else None
            parent_id = emp.get("parent_employee_id") or emp.get("pc_id")
            parent = person_by_id.get(str(parent_id)) if parent_id else None
            parent_label = f"{parent.get('name')} ({parent_id})" if parent else _value(parent_id, "—")
            position = emp.get("designation") or {
                "FS": "Field Supervisor",
                "LS": "Local Supervisor",
                "LSS": "Local Supervisor Scanning",
            }[role]
            ws_pc.append([
                _value(emp.get("employee_id")),
                _value(location or emp.get("area")),
                _value(emp.get("name")),
                parent_label,
                position,
            ] + markers)
            r = ws_pc.max_row
            first = 6
            last = 5 + len(days)
            range_ref = f"{get_column_letter(first)}{r}:{get_column_letter(last)}{r}"
            formulas = [
                f'=COUNTIF({range_ref},"P")+COUNTIF({range_ref},"L")+COUNTIF({range_ref},"M")',
                f'=COUNTIF({range_ref},"A")+COUNTIF({range_ref},"AL")',
                f'=COUNTIF({range_ref},"HD")', f'=COUNTIF({range_ref},"WO")',
                f'=COUNTIF({range_ref},"H")', f'=COUNTIF({range_ref},"L")', f'=COUNTIF({range_ref},"AL")',
            ]
            for offset, formula in enumerate(formulas, start=last + 1):
                ws_pc.cell(r, offset, formula)
            for c, marker in enumerate(markers, start=6):
                cell = ws_pc.cell(r, c)
                cell.fill = MARKER_FILLS.get(marker, PatternFill())
                cell.alignment = Alignment(horizontal="center")
                validation.add(cell)
        _style_sheet(ws_pc)
        ws_pc.freeze_panes = "F2"
        ws_pc.auto_filter.ref = ws_pc.dimensions
        ws_pc.column_dimensions["A"].width = 14
        ws_pc.column_dimensions["B"].width = 24
        ws_pc.column_dimensions["C"].width = 26
        ws_pc.column_dimensions["D"].width = 32
        ws_pc.column_dimensions["E"].width = 28
        for col_idx in range(6, 6 + len(days)):
            ws_pc.column_dimensions[get_column_letter(col_idx)].width = 13
        for col_idx in range(6 + len(days), 13 + len(days)):
            ws_pc.column_dimensions[get_column_letter(col_idx)].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
